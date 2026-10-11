#!/usr/bin/env python3
"""Customer advertising album from REAL, unchanged ARCHIPRO UZ lesson images.

All image files are read directly from app/courses/visual_assets. Each printed
image retains the exact original bytes and has a visible lesson ID.
No AI-generation, substitute illustrations or online stock images.
"""
from __future__ import annotations
from pathlib import Path
import csv
import os
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Inches, Pt as PPTPt
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
ASSET = ROOT / "app" / "courses" / "visual_assets"
OUT = ROOT / "marketing_output"
OUT.mkdir(exist_ok=True)
PDF = OUT / "ARCHIPRO_UZ_10_sahifali_ASL_DARS_ESKIZLARI.pdf"
PPT = OUT / "ARCHIPRO_UZ_10_slaydli_ASL_DARS_ESKIZLARI.pptx"

# Topics selected for professional usefulness and attractiveness.
# Each image code points to its REAL existing, unmodified lesson JPEG.
PAGES = [
    [
        ("CHIZMANI O'QISHNI O'RGANING",
         "Plan, kesim va fasad: bitta loyihani turli tomondan tushunish.",
         ["L15_M05_004","L10_M09_003","L10_M07_002"],
         "Plan + fasad + kesim bir-biriga mos bo'lishini ko'ring."),
        ("FASAD VA KOMPOZITSIYA",
         "Kirish aksenti, proporsiya va derazalar ritmini o'zingiz tanlang.",
         ["L10_M07_003","L10_M07_004","L10_M07_005"],
         "Ko'rkam tashqi ko'rinish sababini chizmadan tushuntiring.")
    ],
    [
        ("XONALAR PLANIROVKASI",
         "Xonalarni shunchaki chizmang: qulay yashashni loyihalang.",
         ["L10_M04_002","L10_M04_001","L10_M04_003"],
         "Yotoqxona, mehmonxona va o'tish yo'laklarini tahlil qiling."),
        ("ERGONOMIKA VA OPTIMIZATSIYA",
         "Har bir kvadrat metr nima uchun kerakligini biling.",
         ["L10_M05_002","L10_M05_003","L10_M05_004"],
         "To'g'ri va noto'g'ri yechimni yonma-yon solishtiring.")
    ],
    [
        ("UCHASTKADAN GENPLANGACHA",
         "Quyosh, yo'l, kirish va hovli zonalarini bir joyga jamlang.",
         ["L10_M10_003","L10_M02_002","L10_M10_005"],
         "Yer uchastkasidan foydali va chiroyli bosh reja oling."),
        ("HOVLI VA MAHALLIY TURMUSH",
         "O'zbekiston sharoitidagi uy: ayvon, kirish va shaxsiy hudud.",
         ["L14_M01_003","L10_M10_002","L14_M01_002"],
         "Ichki hovli, yorug'lik va maxfiylik muvozanatini toping.")
    ],
    [
        ("TOMNING HAQIQIY UZELLARI",
         "Tom qoplamasidan suv o'tkazmaslikkacha - birma-bir ko'ring.",
         ["L10_M08_003","L10_M08_002","L10_M08_004"],
         "Qirra, vodiy va karniz tutashuvlarini eskiz orqali o'rganing."),
        ("KONSTRUKSIYA VA HIMOYA",
         "Poydevor, devor, issiqlik va seysmik fikrlash asoslari.",
         ["L08_M01_004","L08_M07_002","L08_M08_003"],
         "Arxitektura qarori konstruksiyaga qanday ta'sir qilishini biling.")
    ],
    [
        ("ARCHICADDA MODEL YARATISH",
         "Oddiy chizmadan o'lchamli 3D makongacha bo'lgan amaliy yo'l.",
         ["L10_M06_003","L10_M06_004","L10_M06_005"],
         "Devor, deraza, tom va qavatlar bitta modelda ishlaydi."),
        ("BIM FIKRLASH",
         "Geometriya bilan birga element ma'lumoti va tekshiruv.",
         ["L12_M01_005","L12_M03_003","L11_M08_003"],
         "Modeldan o'lcham va jadval ma'lumotini chiqarish.")
    ],
    [
        ("BIR QAVATLI UY: BOSHDAN OXIRIGACHA",
         "Eskiz, kirish, fasad, tom va o'lchamlarning bog'lanishi.",
         ["L15_M05_004","L10_M07_003","L10_M04_002"],
         "Yakuniy ko'rinishning ortidagi barcha chizmalarni ko'ring."),
        ("KESIM, FASAD VA LOYIHA SIFATI",
         "Ochiq qolgan savollarga o'lcham va kesim bilan javob bering.",
         ["L10_M09_003","L10_M09_004","L10_M11_002"],
         "Bitta bino - bir-biriga mos loyiha hujjatlari.")
    ],
    [
        ("IKKI QAVATLI UY KONSEPSIYASI",
         "Hajm, xonalar va qavatlar orasidagi munosabat.",
         ["L11_M01_004","L11_M01_002","L11_M07_002"],
         "Birinchi va ikkinchi qavat bir xil loyihada ishlasin."),
        ("ZINAPOYA VA VERTIKAL ALOQA",
         "Zinapoya shunchaki detal emas, makonning asosiy qismi.",
         ["L11_M03_004","L11_M03_002","L11_M04_002"],
         "Plan, kesim va 3D yechim o'zaro muvofiq bo'lsin.")
    ],
    [
        ("BIM VA SPETSIFIKATSIYA",
         "Element nomi, soni, parametri - barchasi tizimli.",
         ["L12_M01_005","L12_M03_003","L12_M04_004"],
         "BIM modeli bilan ishlovchi loyihachi uslubini o'rganing."),
        ("PROFESSIONAL ALBOM TAYYORLASH",
         "Chizma varag'i, masshtab, sarlavha va ko'rgazmali tartib.",
         ["L15_M05_004","L12_M05_003","L12_M05_004"],
         "Loyihani tushunarli, izchil hujjatlar to'plamiga aylantiring.")
    ],
    [
        ("HAQIQIY MIJOZ BILAN ISHLASH",
         "Talabni tinglash, variant tayyorlash va qarorni himoya qilish.",
         ["L13_M05_003","L13_M04_004","L13_M08_002"],
         "Bir fikrni loyiha ko'rinishida tushuntirishni mashq qiling."),
        ("YAKUNIY LOYIHA VA PORTFOLIO",
         "Eskiz, plan, fasad, BIM va taqdimot - bitta yo'lda.",
         ["L15_M05_004","L15_M05_002","L15_M07_004"],
         "O'rganganlaringizni yakuniy loyiha ishida birlashtiring.")
    ],
]
assert len(PAGES)==9
HERO=["L15_M05_004", "L10_M08_003", "L13_M05_003"]
ALL_CODES=set(HERO+[code for page in PAGES for sec in page for code in sec[2]])

# Verify original source files before publishing anything.
missing=[c for c in ALL_CODES if not (ASSET/("sketch-"+c.lower()+".jpg")).is_file()]
if missing:
    raise SystemExit("ERROR: Actual lesson images missing; refusing to substitute: "+", ".join(missing))
for c in ALL_CODES:
    p=ASSET/("sketch-"+c.lower()+".jpg")
    with Image.open(p) as im: im.verify()
print("Exact existing GitHub lesson images verified:",len(ALL_CODES),flush=True)
DARK=HexColor("#0D1B2A"); NAVY=HexColor("#162E44")
GOLD=HexColor("#CE9A4A"); BG=HexColor("#FBF8F2")
LIGHT=HexColor("#EBE7DD"); MUTED=HexColor("#657586"); WHITE=HexColor("#FFFFFF")
W,H=595.28,841.89
FONT="/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONTB="/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
if not Path(FONT).exists():
    FONT="C:/Windows/Fonts/arial.ttf"; FONTB="C:/Windows/Fonts/arialbd.ttf"
if not Path(FONT).exists(): raise SystemExit("Missing font")
pdfmetrics.registerFont(TTFont("F",FONT));pdfmetrics.registerFont(TTFont("FB",FONTB))
cv=canvas.Canvas(str(PDF),pagesize=(W,H),pageCompression=1)
cv.setTitle("ARCHIPRO UZ | 10 sahifa | Darsdagi asl eskizlar")

def T(s,x,y,size=11,color=DARK,bold=False):
    cv.setFont("FB" if bold else "F",size);cv.setFillColor(color);cv.drawString(x,y,s)

def fitimage(path,x,y,w,h):
    """Draw original JPEG completely without cropping/replacing/editing pixels."""
    with Image.open(path) as im: iw,ih=im.size
    k=min(w/iw,h/ih)
    dw,dh=iw*k,ih*k
    cv.setFillColor(WHITE);cv.rect(x,y,w,h,fill=1,stroke=0)
    cv.drawImage(str(path),x+(w-dw)/2,y+(h-dh)/2,dw,dh,mask="auto")

def card(code,x,y,w,h,hero=False):
    p=ASSET/("sketch-"+code.lower()+".jpg")
    cv.setFillColor(WHITE);cv.roundRect(x,y,w,h,7,stroke=0,fill=1)
    fitimage(p,x+3,y+20,w-6,h-23)
    cv.setFillColor(LIGHT);cv.rect(x+3,y+3,w-6,15,fill=1,stroke=0)
    T("ASL DARS ESKIZI  •  "+code,x+9,y+7,7,DARK,True)

def base(page):
    cv.setFillColor(BG);cv.rect(0,0,W,H,fill=1,stroke=0)
    cv.setFillColor(DARK);cv.rect(0,H-56,W,56,fill=1,stroke=0)
    T("ARCHIPRO UZ",29,H-35,16,WHITE,True)
    T("DARS ICHIDAGI ASL ESKIZLAR",338,H-33,8,GOLD,True)
    cv.setStrokeColor(GOLD);cv.setLineWidth(1);cv.line(29,45,W-29,45)
    T("ELEKTRON ARXITEKTOR USTOZI",30,29,7,MUTED)
    T("%02d / 10"%page,W-84,28,10,DARK,True)

def topic(sec,top,i):
    name,sub,codes,learning=sec
    cv.setFillColor(GOLD);cv.roundRect(29,top-34,29,28,5,stroke=0,fill=1)
    T(str(i).zfill(2),35,top-25,12,DARK,True)
    T(name,67,top-21,14.1,NAVY,True)
    T(sub,67,top-39,8.5,MUTED)
    w=335; rw=163
    card(codes[0],29,top-292,w,243,True)
    card(codes[1],375,top-166,rw,117)
    card(codes[2],375,top-292,rw,117)
    cv.setFillColor(NAVY)
    cv.roundRect(29,top-329,509,29,5,fill=1,stroke=0)
    T("AMALIY NATIJA  •  "+learning,40,top-317,8,WHITE)
    cv.setStrokeColor(LIGHT);cv.line(29,top-342,W-29,top-342)

# Cover: all 3 hero pictures are authentic lesson images.
cv.setFillColor(DARK);cv.rect(0,0,W,H,stroke=0,fill=1)
cv.setFillColor(GOLD);cv.rect(0,H-13,W,6,fill=1,stroke=0)
T("ARCHIPRO UZ",34,H-51,21,WHITE,True)
T("ELEKTRON ARXITEKTOR USTOZI",35,H-73,10,GOLD)
T("ARXITEKTURANI",34,H-122,29,WHITE,True)
T("NOLDAN O'RGANING.",34,H-158,29,GOLD,True)
T("Bu albomdagi har bir rasm - dasturning haqiqiy darsidan.",35,H-184,10,WHITE)
T("Uyingizni tasavvur qiling. Chizmani o'qing. Loyihaga aylantiring.",35,H-202,9,HexColor("#BCCAD7"))
card(HERO[0],29,350,537,263,True)
card(HERO[1],29,178,262,162)
card(HERO[2],304,178,262,162)
for j,(num,lbl) in enumerate([("16","BOSQICH"),("137","MODUL"),("680","DARS"),("2041","TEST SAVOLI")]):
    x=30+j*136
    cv.setFillColor(NAVY);cv.roundRect(x,91,125,69,6,fill=1,stroke=0)
    T(num,x+10,125,22,GOLD,True);T(lbl,x+11,107,9,WHITE,True)
T("SOTUV ALBOMI    •    FAQAT DARS ICHIDAGI ASL TASVIRLAR",31,51,8,WHITE)
T("01 / 10",W-82,24,10,GOLD,True)
cv.showPage()
for n,page in enumerate(PAGES,2):
    base(n)
    topic(page[0],777,1)
    topic(page[1],408,2)
    cv.showPage()
cv.save()

# Editable widescreen presentation, using exactly the same original lesson JPEGs.
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
def pc(hexv): return RGBColor.from_string(hexv.lstrip("#"))
def rect(sl,x,y,w,h,color):
    sh=sl.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=pc(color)
    sh.line.fill.background()
def txt(sl,text,x,y,w,h,size,color,bold=False):
    tb=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=tb.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_top=0
    pr=tf.paragraphs[0];pr.text=text
    pr.font.name="DejaVu Sans";pr.font.size=PPTPt(size)
    pr.font.bold=bold;pr.font.color.rgb=pc(color)
def pic(sl,code,x,y,w,h):
    img=ASSET/("sketch-"+code.lower()+".jpg")
    with Image.open(img) as im:iw,ih=im.size
    k=min(w/iw,h/ih);dw,dh=iw*k,ih*k
    rect(sl,x,y,w,h,"FFFFFF")
    sl.shapes.add_picture(str(img), Inches(x+(w-dw)/2), Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh))
    txt(sl,"DARS "+code,x+0.10,y+h-0.23,w-0.1,0.19,8,"17314A",True)
s=prs.slides.add_slide(BLANK);rect(s,0,0,13.333,7.5,"0D1B2A")
txt(s,"ARCHIPRO UZ",0.6,0.4,6,0.5,23,"FFFFFF",True)
txt(s,"ARXITEKTURANI NOLDAN O'RGANING",0.6,1.1,11.8,0.8,34,"E9C384",True)
txt(s,"16 bosqich   |   137 modul   |   680 dars   |   2041 test",0.6,2.06,11.8,0.5,16,"FFFFFF")
pic(s,HERO[0],0.6,2.75,6.1,4.1);pic(s,HERO[1],6.85,2.75,5.85,1.9);pic(s,HERO[2],6.85,4.85,5.85,2)
for page_num,page in enumerate(PAGES,2):
    s=prs.slides.add_slide(BLANK);rect(s,0,0,13.333,7.5,"F9F6EF")
    rect(s,0,0,13.333,0.65,"0D1B2A")
    txt(s,"ARCHIPRO UZ  •  ASL DARS ESKIZLARI",0.45,0.19,9,0.38,17,"FFFFFF",True)
    txt(s,f"{page_num:02d} / 10",11.4,0.17,1.5,0.38,17,"E9C384",True)
    for j,sec in enumerate(page):
        y=0.85+j*3.28
        title,sub,codes,learning=sec
        txt(s,title,0.42,y,11,0.38,19,"162E44",True)
        txt(s,sub,0.42,y+0.4,12,0.3,10,"657586")
        pic(s,codes[0],0.45,y+0.78,6.05,2.31)
        pic(s,codes[1],6.67,y+0.78,3.04,2.31)
        pic(s,codes[2],9.88,y+0.78,3.04,2.31)
prs.save(PPT)
print("READY",PDF,PDF.stat().st_size,PPT,PPT.stat().st_size,flush=True)
